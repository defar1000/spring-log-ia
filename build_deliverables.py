from pathlib import Path
import json, os, re, shutil, textwrap, zipfile, statistics
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference, LineChart
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import DataBarRule

ROOT=Path(__file__).resolve().parent
DATA=ROOT/'data'; OUT=ROOT/'outputs'; CHART=ROOT/'charts'
NAVY='16324F'; TEAL='0F8B8D'; LIGHT='EAF2F8'; PALE='F4F7FA'; TEXT='243447'; MUTED='5B6573'; AMBER='E0A126'; RED='B23A48'; GREEN='2E7D63'; WHITE='FFFFFF'
SEV_COLORS={'INFO':'0F8B8D','WARN':'E0A126','ERROR':'D95D39','CRITICAL':'9B2226'}

clean=pd.read_csv(DATA/'payment_service_logs_500_clean.csv',dtype=str,keep_default_na=False)
raw=pd.read_csv(DATA/'payment_service_logs_500.csv',dtype=str,keep_default_na=False)
eda=json.load(open(OUT/'eda_summary.json',encoding='utf-8'))
metrics=json.load(open(OUT/'model_metrics.json',encoding='utf-8'))
model_cmp=pd.read_csv(OUT/'model_comparison.csv')
report=pd.read_csv(OUT/'classification_report.csv',index_col=0)
audit=pd.read_csv(OUT/'audit_sample_25.csv',keep_default_na=False)

# --------------------- DOCX design system ---------------------
def shade_cell(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); mar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_table_borders(table, color='D5DEE8', size='4'):
    tblPr=table._tbl.tblPr; borders=tblPr.first_child_found_in('w:tblBorders')
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; element=borders.find(qn(tag))
        if element is None: element=OxmlElement(tag); borders.append(element)
        element.set(qn('w:val'),'single'); element.set(qn('w:sz'),size); element.set(qn('w:space'),'0'); element.set(qn('w:color'),color)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def set_keep(paragraph, keep=True):
    pPr=paragraph._p.get_or_add_pPr(); el=OxmlElement('w:keepNext' if keep else 'w:keepLines'); pPr.append(el)

def add_field(run, field):
    fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=field
    fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end')
    run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)

def set_cell_text(cell, text, bold=False, color=TEXT, size=9, align=None):
    cell.text=''
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
    if align is not None: p.alignment=align
    run=p.add_run(str(text)); run.bold=bold; run.font.name='Aptos'; run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
    return p

def add_doc_table(doc, headers, rows, widths=None, font_size=8.5, header_fill=NAVY, aligns=None, repeat=True):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    if widths:
        for i,w in enumerate(widths):
            for cell in table.columns[i].cells: cell.width=Inches(w)
    for j,h in enumerate(headers):
        c=table.rows[0].cells[j]; shade_cell(c,header_fill); set_cell_text(c,h,True,WHITE,font_size)
        if widths: c.width=Inches(widths[j])
    if repeat: set_repeat_table_header(table.rows[0])
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for j,val in enumerate(row):
            fill=WHITE if ri%2==0 else PALE
            shade_cell(cells[j],fill)
            align=None
            if aligns and j<len(aligns) and aligns[j]=='center': align=WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[j],val,False,TEXT,font_size,align)
            if widths: cells[j].width=Inches(widths[j])
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return table

def add_callout(doc, title, body, fill=LIGHT, accent=TEAL):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); shade_cell(c,fill); set_cell_margins(c,150,190,150,190)
    c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(4)
    r=p.add_run(title); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(accent)
    p2=c.add_paragraph(); p2.paragraph_format.space_after=Pt(0); p2.paragraph_format.line_spacing=1.08
    r=p2.add_run(body); r.font.name='Aptos'; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(TEXT)
    set_table_borders(t,color=fill,size='0'); doc.add_paragraph().paragraph_format.space_after=Pt(1)

def add_p(doc,text='',bold_prefix=None,italic=False,align=None,after=5):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.08; p.paragraph_format.keep_together=True
    if align is not None: p.alignment=align
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
        p.add_run(text[len(bold_prefix):])
    else:
        r=p.add_run(text); r.italic=italic
    for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor.from_string(TEXT)
    return p

def add_bullets(doc, items, numbered=False):
    for item in items:
        p=doc.add_paragraph(style='List Number' if numbered else 'List Bullet')
        p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.15)
        p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.08; p.paragraph_format.keep_together=True
        r=p.add_run(str(item)); r.font.name='Aptos'; r.font.size=Pt(9.3); r.font.color.rgb=RGBColor.from_string(TEXT)

def add_heading(doc, text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.space_before=Pt(9 if level==1 else 6); p.paragraph_format.space_after=Pt(4)
    p.add_run(text)
    set_keep(p)
    return p

def add_chart(doc, path, width=6.2, caption=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
    p.add_run().add_picture(str(path),width=Inches(width))
    if caption:
        cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after=Pt(5)
        r=cp.add_run(caption); r.italic=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)

def add_chart_grid(doc, items):
    for start in range(0,len(items),2):
        subset=items[start:start+2]
        t=doc.add_table(rows=1,cols=len(subset)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
        for j,(path,caption) in enumerate(subset):
            c=t.cell(0,j); c.text=''; c.width=Inches(3.25); set_cell_margins(c,60,60,60,60)
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
            p.add_run().add_picture(str(path),width=Inches(3.05))
            q=c.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(1)
            r=q.add_run(caption); r.italic=True; r.font.name='Aptos'; r.font.size=Pt(7.5); r.font.color.rgb=RGBColor.from_string(MUTED)
        set_table_borders(t,color=WHITE,size='0')
        doc.add_paragraph().paragraph_format.space_after=Pt(0)

def base_doc(code,title,subtitle):
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    sec.header_distance=Inches(.28); sec.footer_distance=Inches(.3)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor.from_string(TEXT)
    styles['Normal'].paragraph_format.space_after=Pt(5)
    for nm,size,color in [('Title',26,NAVY),('Subtitle',11,MUTED),('Heading 1',15,NAVY),('Heading 2',11,TEAL),('Heading 3',10,NAVY)]:
        st=styles[nm]; st.font.name='Aptos Display' if nm in ('Title','Heading 1','Heading 2') else 'Aptos'; st.font.size=Pt(size); st.font.bold=nm!='Subtitle'; st.font.color.rgb=RGBColor.from_string(color)
    h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; h.paragraph_format.space_after=Pt(2)
    r=h.add_run('DIPLOMADO EN INTELIGENCIA ARTIFICIAL  |  OPCIÓN DE GRADO'); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(7.5); r.font.color.rgb=RGBColor.from_string(TEAL)
    pPr=h._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'16'); bottom.set(qn('w:space'),'5'); bottom.set(qn('w:color'),TEAL); pbdr.append(bottom); pPr.append(pbdr)
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; f.paragraph_format.space_before=Pt(2)
    r=f.add_run('SPRINGLOG-IA  |  PROYECTO FINAL INTEGRADO  •  '); r.font.name='Aptos'; r.font.size=Pt(7.5); r.font.color.rgb=RGBColor.from_string(MUTED)
    r=f.add_run('Pág. '); r.font.size=Pt(7.5); r.font.color.rgb=RGBColor.from_string(MUTED); add_field(r,'PAGE')
    r=f.add_run(' de '); r.font.size=Pt(7.5); r.font.color.rgb=RGBColor.from_string(MUTED); add_field(r,'NUMPAGES')
    # Cover block consistent across all technical documents.
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(2)
    r=p.add_run('SPRINGLOG-IA'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string(TEAL)
    p=doc.add_paragraph(style='Title'); p.paragraph_format.space_after=Pt(2); p.add_run(title)
    p=doc.add_paragraph(style='Subtitle'); p.paragraph_format.space_after=Pt(8); p.add_run(subtitle)
    meta=doc.add_table(rows=2,cols=4); meta.alignment=WD_TABLE_ALIGNMENT.CENTER; meta.autofit=False
    headers=[('ENTREGABLE',code),('ESTUDIANTE','David Felipe Rios Echeverri'),('ÁREA','Desarrollo Backend / Ingeniería de Software'),('FECHA','Octubre de 2026')]
    for idx,(k,v) in enumerate(headers):
        c=meta.cell(idx//2,(idx%2)*2); d=meta.cell(idx//2,(idx%2)*2+1)
        shade_cell(c,NAVY); set_cell_text(c,k,True,WHITE,7.5)
        shade_cell(d,PALE); set_cell_text(d,v,False,TEXT,8.3)
        c.width=Inches(1.15); d.width=Inches(2.15)
    set_table_borders(meta,color='D5DEE8',size='4')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    doc.core_properties.title=f'{code} - {title}'; doc.core_properties.author='David Felipe Rios Echeverri'; doc.core_properties.subject='SpringLog-IA'
    return doc

def save_doc(doc, filename):
    path=ROOT/filename; doc.save(path); return path

# --------------------- 01 Ficha del problema ---------------------
doc=base_doc('01','Ficha del problema','Definición del problema, datos, objetivo y criterios de éxito')
add_callout(doc,'Resumen del proyecto','SpringLog-IA propone un clasificador asistivo para priorizar logs de un microservicio Java/Spring Boot. El prototipo actual se enfoca en predecir severidad; el diagnóstico de causa raíz, la detección de anomalías y la integración en observabilidad quedan como extensiones futuras.')
add_heading(doc,'1. Contexto',1)
add_p(doc,'En arquitecturas backend basadas en Java y Spring Boot, microservicios y sus dependencias producen eventos operativos que ayudan a detectar errores, latencia, problemas de conexión y fallos de integración. Durante un incidente, el equipo necesita distinguir rápidamente el tráfico rutinario de señales que requieren atención, correlacionando la información entre instancias.')
add_heading(doc,'2. Planteamiento del problema',1)
add_p(doc,'La revisión manual de líneas de log y stack traces puede consumir tiempo valioso de soporte. Los eventos INFO y WARN conviven con excepciones de mayor severidad; filtros por palabras clave pueden perder contexto técnico y dificultar la priorización. Las estimaciones iniciales de 45 minutos por diagnóstico y 40–60 % de tiempo dedicado a filtrar logs se consideran hipótesis operativas de partida: no fueron medidas en producción para este proyecto.')
add_heading(doc,'3. Datos disponibles',1)
add_doc_table(doc,['Atributo','Descripción'],[
('Origen','Dataset sintético para un único microservicio payment-service; 500 registros.'),
('Periodo','2026-10-01 00:29:10 a 2026-10-05 23:58:52.'),
('Instancias','instance-prod-east-01, instance-prod-east-02, instance-prod-west-01.'),
('Campos originales','timestamp, app_name, instance_id, severity, thread, logger_class, exception_type, log_message.'),
('Severidad','INFO 300 (60 %); WARN 100 (20 %); ERROR 75 (15 %); CRITICAL 25 (5 %).'),
('Naturaleza','Datos sintéticos; no representan telemetría real ni estiman frecuencia real de incidentes.')],[1.5,5.0],8.4)
add_heading(doc,'4. Objetivo con IA',1)
add_p(doc,'Diseñar y evaluar un prototipo de clasificación automática de severidad de logs backend, usando características textuales de la clase emisora, el tipo de excepción y el mensaje, para apoyar la priorización humana de incidentes.')
add_bullets(doc,[
'Validar y preparar el CSV, normalizando marcadores explícitos de ausencia y sanitizando datos de identificación o secretos antes de generar características.',
'Entrenar y comparar clasificadores de texto con una partición estratificada de entrenamiento y prueba.',
'Priorizar el recall de las categorías ERROR y CRITICAL y documentar los errores de WARN/INFO.',
'Definir un proceso futuro de prueba con logs reales anonimizados, revisión humana e integración con herramientas de observabilidad.'
])
add_heading(doc,'5. Restricciones técnicas, éticas y legales',1)
add_bullets(doc,[
'La muestra es sintética y pequeña; no permite afirmar generalización a producción ni cuantificar el MTTR real.',
'Los mensajes pueden contener identificadores, tokens, correos, credenciales u otros datos sensibles. La sanitización debe ejecutarse antes de cualquier envío a servicios externos.',
'La clasificación será asistiva. Un ingeniero debe validar diagnósticos y decidir acciones operativas; el modelo no ejecuta remediaciones.',
'El prototipo local se construye con Python y scikit-learn. El tiempo de inferencia depende del hardware y del tamaño del lote.'
])
add_heading(doc,'6. Criterios de éxito y estado de la evidencia',1)
add_doc_table(doc,['Indicador','Criterio propuesto','Estado actual'],[
('Clasificación general','Accuracy y F1 macro ≥ 0,85 en prueba reservada.','Cumplido en la partición sintética; requiere validación externa.'),
('Prioridad crítica','Recall de CRITICAL ≥ 0,90 en prueba.','5/5 CRITICAL detectados en la partición de prueba; soporte pequeño.'),
('Latencia del prototipo','P95 de inferencia < 50 ms por registro en entorno de prueba.',f"Medición local: P95 aproximado {metrics['latency']['p95_ms']:.1f} ms; depende del entorno."),
('Reducción del ruido','Medir reducción de INFO irrelevantes durante incidentes.','Pendiente de piloto; no medida en este proyecto.'),
('Impacto en MTTR','Comparar tiempo de diagnóstico manual vs asistido.','Pendiente; cifras de ahorro son escenarios de negocio, no resultado observado.')],[1.4,2.55,2.55],8.1)
add_callout(doc,'Alcance de la primera versión','Se evalúa clasificación de severidad. No se presenta como implementada la predicción de fallas, la detección independiente de anomalías, el diagnóstico definitivo de causa raíz ni la integración en tiempo real.')
save_doc(doc,'01_FICHA_DEL_PROBLEMA.docx')

# --------------------- 02 IA generativa ---------------------
doc=base_doc('02','Solución con IA generativa','Prompts especializados, ejemplos, capacidades y limitaciones')
add_callout(doc,'Propósito','Explorar cómo la IA generativa puede estructurar logs, sugerir hipótesis de diagnóstico, ayudar a sanitizar datos y resumir patrones. Sus salidas se consideran asistencia para un ingeniero, no evidencia por sí solas de la causa raíz.')
add_heading(doc,'1. Contexto de la exploración',1)
add_p(doc,'La exploración se realiza sobre un conjunto sintético de 500 registros de payment-service. Los ejemplos de entrada y salida de este informe son ejemplos reproducibles de diseño de interacción, no una prueba comparativa de modelos generativos ni un resultado validado sobre telemetría de producción.')
add_heading(doc,'2. Prompts especializados',1)
prompts=[
('Prompt 1 — Extracción estructurada y clasificación','Actúa como ingeniero SRE sénior experto en Java y Spring Boot 3. Analiza el evento recibido y devuelve únicamente JSON válido con timestamp, app_name, instance_id, thread, logger_class, exception_type, severity_original, severity_sugerida, resumen_tecnico y confianza. Conserva la severidad original. Usa null o NO_APLICA si falta un campo. No inventes información. Separa los hechos observados de cualquier inferencia. Log: [PEGAR LOG SANITIZADO].'),
('Prompt 2 — Diagnóstico de causa raíz','Actúa como arquitecto backend experto en Spring Data JPA, Hibernate y HikariCP. A partir de la traza, separa hechos, hipótesis e información faltante. Entrega causa raíz probable, evidencia textual, impacto potencial, pruebas de diagnóstico priorizadas y métricas para verificar la hipótesis. No propongas cambios de configuración sin evidencia suficiente. Traza: [PEGAR STACK TRACE SANITIZADO].'),
('Prompt 3 — Sanitización de datos sensibles','Actúa como responsable de seguridad. Detecta y enmascara JWT/Bearer, tokens, contraseñas, claves API, correos, identificadores personales y números de tarjeta. Sustituye los valores por marcadores [REDACTED_*], conserva el contexto técnico no sensible y devuelve la línea sanitizada y la lista de tipos de dato ocultados. No reproduzcas los valores originales. Log: [PEGAR LOG].'),
('Prompt 4 — Resumen operativo','Actúa como líder técnico de payment-service. Resume los registros suministrados por severidad e identifica tipos de excepción y señales operativas recurrentes. Diferencia conteos del dataset, hipótesis de impacto y recomendaciones. No sumes categorías solapadas como si fueran incidentes únicos. Declara que la muestra es sintética. Datos: [PEGAR TABLA O LOGS].')]
for title,body in prompts:
    if title.startswith('Prompt 3'):
        doc.add_page_break()
    add_heading(doc,title,2); add_callout(doc,'Prompt',body,fill=PALE,accent=NAVY)
add_heading(doc,'3. Ejemplos de entrada y salida',1)
add_heading(doc,'3.1. Caso de pool de conexiones',2)
add_p(doc,'Entrada sintética: “TransactionRepository: Connection is not available; HikariPool timeout after 30000 ms; activeConnections=20; maxPoolSize=20”.')
add_p(doc,'Salida ilustrativa: severidad sugerida CRITICAL si coincide con la etiqueta y el contexto de operación; síntoma observado: no se consiguió una conexión dentro del tiempo límite; hipótesis por verificar: saturación del pool, consultas lentas o transacciones largas; acción inicial: revisar métricas active/idle/pending/timeouts y límites de la base de datos antes de cambiar parámetros.')
add_heading(doc,'3.2. Caso de sanitización',2)
add_p(doc,'Entrada ficticia: “Payment failed userEmail=ana@example.test token=abc123secret customerId=9281”.')
add_p(doc,'Salida esperada: “Payment failed userEmail=[REDACTED_EMAIL] token=[REDACTED] customerId=[REDACTED_ID]”.')
add_heading(doc,'3.3. Resumen de la muestra',2)
add_doc_table(doc,['Severidad','Registros','Porcentaje'],[[s,int(eda['severity_counts'][s]),f"{eda['severity_percent'][s]:.0f} %"] for s in ['INFO','WARN','ERROR','CRITICAL']],[2.1,1.6,1.5],8.6)
add_p(doc,'La proporción de severidades fue definida al generar el dataset; no es una estimación de la frecuencia de eventos de producción.')
add_heading(doc,'4. Capacidades observadas y límites',1)
add_doc_table(doc,['Capacidad','Utilidad','Limitación'],[
('Estructuración','Normaliza campos y genera JSON tabular.','Puede inventar o inferir campos si el prompt no restringe la salida.'),
('Diagnóstico asistido','Propone hipótesis y pasos de verificación.','Un log aislado no demuestra causa raíz.'),
('Resumen','Sintetiza frecuencias y patrones.','Categorías por palabras clave pueden solaparse.'),
('Sanitización','Ayuda a revisar la política de enmascaramiento.','Un LLM no debe ser el único control; se requiere filtro determinista y pruebas.'),
('Escalabilidad','Puede asistir análisis complejos.','Costo, latencia, límites de API y privacidad impiden asumir que procesará cada log en tiempo real.')],[1.25,2.5,2.9],8.0)
add_heading(doc,'5. Próximos pasos',1)
add_bullets(doc,['Aplicar sanitización determinista antes de cualquier servicio externo.','Versionar prompts y conservar ejemplos de evaluación con datos sintéticos.','Comparar resultados generativos con el clasificador supervisado usando el mismo conjunto de casos.','Requerir revisión humana para CRITICAL, confianza baja o recomendaciones que impliquen cambios operativos.'])
save_doc(doc,'02_IA_GENERATIVA.docx')

# --------------------- 03 EDA report ---------------------
doc=base_doc('03','Análisis de datos exploratorio (EDA)','Calidad de datos, distribución, patrones y dataset preparado')
add_callout(doc,'Resumen ejecutivo',f"Se analizaron {len(clean)} registros sintéticos de payment-service entre {eda['timestamp_start']} y {eda['timestamp_end']}. No se encontraron duplicados exactos en los campos originales. El CSV contenía {eda['source_exception_none_sentinel']} marcadores literales “None” en exception_type; no eran nulos CSV, por lo que se normalizaron a NO_APLICA. La distribución de clases está desbalanceada: INFO 60 % y CRITICAL 5 %.")
add_heading(doc,'1. Inventario y calidad de datos',1)
add_doc_table(doc,['Elemento','Resultado','Interpretación'],[
('Registros',str(eda['n_rows']),'500 filas en un microservicio sintético.'),
('Atributos fuente',str(eda['raw_columns']),'Ocho columnas originales con timestamp, contexto y mensaje.'),
('Atributos limpios',str(eda['clean_columns']),'Se añadieron date, hour y day_name para análisis temporal.'),
('Duplicados exactos',str(eda['exact_duplicates']),'Ninguno detectado en los registros procesados.'),
('Nulos CSV reales','0 después de validación','El literal “None” de exception_type era un marcador textual, no un nulo real.'),
('Marcadores de excepción ausente',str(eda['source_exception_none_sentinel']),'Normalizados a NO_APLICA para evitar confusión al volver a leer el CSV.'),
('Ventana temporal',f"{eda['timestamp_start']} a {eda['timestamp_end']}",'Cinco días de registros sintéticos.')],[1.65,1.65,3.15],8.1)
add_heading(doc,'2. Distribución de severidad',1)
add_chart(doc,CHART/'eda_severidad.png',width=5.8,caption='Figura 1. Distribución de severidades del dataset.')
add_p(doc,'INFO representa 300/500 registros (60 %), WARN 100 (20 %), ERROR 75 (15 %) y CRITICAL 25 (5 %). El diseño deliberado de estas proporciones provoca un desequilibrio de clases que debe considerarse en el entrenamiento y la evaluación.')
add_heading(doc,'3. Distribución por instancia',1)
node_rows=[]
for node in ['instance-prod-east-01','instance-prod-east-02','instance-prod-west-01']:
    subset=clean[clean.instance_id==node]; total=len(subset)
    node_rows.append([node, int((subset.severity=='INFO').sum()), int((subset.severity=='WARN').sum()), int((subset.severity=='ERROR').sum()), int((subset.severity=='CRITICAL').sum()), total, f'{total/len(clean):.1%}'])
add_doc_table(doc,['Instancia','INFO','WARN','ERROR','CRITICAL','Total','%'],node_rows,[1.9,.55,.55,.6,.7,.55,.55],7.7)
add_chart(doc,CHART/'eda_instancias.png',width=5.9,caption='Figura 2. Distribución de severidad por instancia.')
add_heading(doc,'4. Patrones temporales y técnicos',1)
add_chart_grid(doc,[(CHART/'eda_hora.png','Figura 3. Volumen por hora.'),(CHART/'eda_dia.png','Figura 4. Eventos por día y severidad.'),(CHART/'eda_excepciones.png','Figura 5. Excepciones más frecuentes.'),(CHART/'eda_longitud.png','Figura 6. Longitud del mensaje por severidad.')])
add_p(doc,'Entre los tipos de excepción más frecuentes del conjunto figuran MethodArgumentNotValidException (29), HttpServerErrorException$ServiceUnavailable (16), HikariPoolTimeoutException (13) y CannotGetJdbcConnectionException (12). Estos conteos son observaciones de esta muestra sintética, no predicciones de tasa de fallos real.')
doc.add_page_break()
add_heading(doc,'5. Preparación para el modelo',1)
add_bullets(doc,[
'Convertir timestamp a fecha/hora y derivar date, hour y day_name para análisis exploratorio.',
'Convertir el marcador literal None de exception_type a NO_APLICA; volver a comprobar que no existen celdas vacías.',
'Sanitizar JWT/Bearer, tokens/secretos, correos, identificadores y números de tarjeta en los mensajes procesados.',
'Mitigar atajos de etiquetado durante el modelado sustituyendo menciones textuales de niveles (INFO/WARN/ERROR/CRITICAL) por [LEVEL] en las características.',
'Mantener la severidad como variable objetivo; no incluirla dentro del texto de entrada del clasificador.'
])
add_heading(doc,'6. Sesgos y limitaciones del EDA',1)
add_p(doc,'La distribución se definió al crear los datos y los patrones técnicos están muy asociados a severidad. Por tanto, un modelo puede aprender reglas específicas de las plantillas sintéticas y obtener resultados que no se trasladan a logs reales. El conjunto no incluye telemetría de producción, etiquetas auditadas externamente ni datos suficientes para evaluar deriva temporal. La validación con mensajes nuevos y un conjunto etiquetado por expertos sigue siendo un requisito previo a producción.')
add_callout(doc,'Artefactos generados','El paquete incluye el CSV original en data/, la versión limpia payment_service_logs_500_clean.csv, el script 03_ANALISIS_DATOS.py, este reporte y el libro 03_ANALISIS_DATOS.xlsx.')
save_doc(doc,'03_ANALISIS_DATOS.docx')

# --------------------- 04 Model report ---------------------
doc=base_doc('04','Modelo o prototipo de IA','Clasificación supervisada de severidad con TF-IDF y Random Forest')
add_callout(doc,'Resumen técnico',f"Se entrenó y evaluó un pipeline TF-IDF (unigramas/bigramas, máximo 1.000 características) + Random Forest de 100 árboles. En la partición de prueba estratificada de 100 registros obtuvo accuracy {metrics['accuracy']:.1%} y F1 macro {metrics['f1_macro']:.3f}. El artefacto de inferencia se guarda como springlog_random_forest.joblib.")
add_heading(doc,'1. Arquitectura del pipeline',1)
add_doc_table(doc,['Etapa','Implementación','Control o nota'],[
('Ingesta','CSV normalizado de payment-service.','Validación de columnas requeridas y severidades permitidas.'),
('Sanitización','Regex para JWT/Bearer, secretos/tokens, correos, IDs y PAN.','La sanitización implementada es basada en patrones; no garantiza detectar todo secreto posible.'),
('Representación','Concatenación de logger_class + exception_type + log_message.','Los valores de ausencia pasan a NO_APLICA; menciones de severidad se enmascaran como [LEVEL].'),
('Vectorización','TF-IDF con n-gramas (1,2), max_features=1000 y sublinear_tf=True.','El vectorizador se ajusta solo en el conjunto de entrenamiento durante evaluación.'),
('Clasificador','RandomForestClassifier(n_estimators=100, class_weight=balanced, random_state=42).','Se conserva como baseline por continuidad del proyecto; no es el algoritmo de mayor puntuación en esta muestra.')],[1.05,2.45,2.95],8.0)
doc.add_page_break()
add_heading(doc,'2. Proceso de entrenamiento y comparación',1)
add_p(doc,'La evaluación utiliza train_test_split con test_size=0.20, stratify=y y random_state=42: 400 registros para entrenamiento y 100 para prueba. Se comparan cuatro clasificadores sobre la misma representación textual.')
model_rows=[]
for _,r in model_cmp.iterrows():
    model_rows.append([r['Modelo'],f"{r['Accuracy']:.1%}",f"{r['Precision_macro']:.3f}",f"{r['Recall_macro']:.3f}",f"{r['F1_macro']:.3f}"])
add_doc_table(doc,['Modelo','Accuracy','Precisión macro','Recall macro','F1 macro'],model_rows,[2.35,.8,1.15,1.05,1.0],8.0)
add_chart(doc,CHART/'modelo_comparacion.png',width=5.7,caption='Figura 1. Comparación de los cuatro modelos sobre la misma partición.')
add_callout(doc,'Interpretación de la comparación','Regresión logística y Naive Bayes alcanzaron 100 % en esta partición, y el árbol de decisión 98 %. En una muestra sintética con patrones muy regulares, resultados perfectos son una señal para ampliar la validación, no una garantía de generalización. Se mantiene Random Forest como prototipo de referencia, sin afirmar que supere a los demás modelos.')
add_heading(doc,'3. Desempeño del modelo de referencia',1)
class_rows=[]
for lab in ['CRITICAL','ERROR','INFO','WARN']:
    r=report.loc[lab]
    class_rows.append([lab,f"{r['precision']:.2f}",f"{r['recall']:.2f}",f"{r['f1-score']:.2f}",int(r['support'])])
class_rows.append(['Macro promedio',f"{metrics['precision_macro']:.2f}",f"{metrics['recall_macro']:.2f}",f"{metrics['f1_macro']:.2f}",100])
class_rows.append(['Promedio ponderado','0.94','0.93',f"{metrics['f1_weighted']:.2f}",100])
add_doc_table(doc,['Clase','Precisión','Recall','F1','Soporte'],class_rows,[1.8,1.0,1.0,1.0,1.0],8.3)
add_chart(doc,CHART/'modelo_matriz_confusion.png',width=4.8,caption='Figura 2. Filas = clase real; columnas = clase predicha. Orden: CRITICAL, ERROR, INFO, WARN.')
add_heading(doc,'4. Ejemplo de predicción y uso',1)
add_p(doc,'Entrada de demostración (mensaje sintético sanitizado): “com.empresa.payment.repository.TransactionRepository HikariPoolTimeoutException [LEVEL] database connection pool exhausted pool=HikariPool-1 timeoutMs=30000 activeConnections=20 maxPoolSize=20”. El script prepara el texto de entrada del mismo modo que durante entrenamiento y genera una etiqueta predicha. La salida es una clasificación, no un diagnóstico definitivo de causa raíz.')
add_heading(doc,'5. Medición de inferencia local',1)
lat=metrics['latency']
add_p(doc,f"En la ejecución registrada se midieron predicciones individuales en este entorno local sobre 30 registros de prueba: mediana {lat['median_ms']:.1f} ms y percentil 95 {lat['p95_ms']:.1f} ms por registro. Es una medición orientativa, dependiente del hardware y carga; no equivale al MTTR ni a latencia de un servicio de producción.")
add_heading(doc,'6. Limitaciones y trabajo futuro',1)
add_bullets(doc,[
'La muestra sintética y sus plantillas asociadas a severidad pueden exagerar el desempeño fuera de muestra.',
'La clasificación de WARN es el principal punto débil: recall 65 %; 7 de 20 WARN se predijeron como INFO.',
'No se implementa integración online con Logback/Log4j2, Prometheus o Grafana.',
'Antes de producción, medir generalización en plantillas nuevas, revisar datos con expertos, calibrar umbrales y mantener supervisión humana.',
'El pipeline serializado usa como entrada una lista de cadenas combinadas y sanitizadas; el README explica el contrato de inferencia.'
])
save_doc(doc,'04_MODELO_O_PROTOTIPO.docx')

# --------------------- 05 evaluation ---------------------
doc=base_doc('05','Evaluación de resultados y validación','Métricas, casos borde y alcance de las conclusiones')
add_callout(doc,'Resultado principal',f"Random Forest obtuvo {metrics['accuracy']:.1%} de accuracy, F1 macro {metrics['f1_macro']:.3f} y F1 ponderado {metrics['f1_weighted']:.3f} sobre 100 registros de prueba. Detectó 5/5 ejemplos CRITICAL y 15/15 ERROR, mientras que clasificó 7 de los 20 WARN como INFO. La muestra es sintética y pequeña.")
add_heading(doc,'1. Diseño de evaluación',1)
add_doc_table(doc,['Elemento','Configuración'],[
('Dataset total','500 registros sintéticos de payment-service.'),
('Partición','80 % entrenamiento (400) / 20 % prueba (100), estratificada.'),
('Reproducibilidad','random_state=42; TF-IDF se ajusta dentro del pipeline y sobre entrenamiento.'),
('Métrica principal','F1 macro junto con accuracy y métricas por clase, para no ocultar errores en clases minoritarias.'),
('Auditoría auxiliar','Se exportan 25 casos estratificados del test para revisión textual. No equivale a auditoría independiente realizada por un experto de producción.')],[1.7,4.8],8.4)
add_heading(doc,'2. Métricas globales y por clase',1)
add_doc_table(doc,['Indicador','Valor','Interpretación'],[
('Accuracy',f"{metrics['accuracy']:.1%}",'93 de 100 predicciones correctas.'),
('Precision macro',f"{metrics['precision_macro']:.3f}",'Promedio de precisión por clase.'),
('Recall macro',f"{metrics['recall_macro']:.3f}",'Promedio de sensibilidad por clase.'),
('F1 macro',f"{metrics['f1_macro']:.3f}",'Métrica principal sensible a desempeño en clases minoritarias.'),
('F1 ponderado',f"{metrics['f1_weighted']:.3f}",'Promedio ponderado por el soporte de las clases.'),
('CRITICAL recall','100 % (5/5)','Sin omisiones CRITICAL en esta partición; soporte solo cinco casos.'),
('ERROR recall','100 % (15/15)','Todos los ERROR de prueba detectados; no garantiza desempeño fuera de la muestra.'),
('WARN recall','65 % (13/20)','7 de 20 WARN fueron clasificados como INFO.')],[1.55,1.35,3.6],8.1)
add_chart(doc,CHART/'modelo_matriz_confusion.png',width=4.9,caption='Figura 1. Matriz de confusión de Random Forest.')
add_heading(doc,'3. Casos borde',1)
add_p(doc,'El error observado se concentra en WARN que se confunden con INFO. En la matriz, 7 WARN se predicen como INFO; no se observan errores CRITICAL→INFO ni ERROR→INFO en la partición. Esto describe solo el conjunto de prueba y no debe expresarse como garantía de seguridad operacional.')
add_bullets(doc,[
'Posible causa: logs WARN sin excepción explícita pueden compartir vocabulario con mensajes rutinarios; el modelo textual no utiliza durationMs como variable numérica independiente.',
'Mejora técnica: extraer duración, estado HTTP, intentos de reintento y métricas del pool como variables estructuradas.',
'Mejora de validación: mantener un conjunto temporal o por plantilla que no se use en entrenamiento y realizar revisión por ingeniero backend/SRE.'
])
add_heading(doc,'4. Comparación antes/después y efecto operativo',1)
add_p(doc,f"No existe todavía una medición experimental del tiempo de diagnóstico antes y después de utilizar SpringLog-IA. El tiempo de inferencia del modelo (mediana {metrics['latency']['median_ms']:.1f} ms/registro en el entorno de prueba) mide una etapa computacional, no el tiempo de resolución del incidente. Las cifras de reducción del MTTR deben mantenerse como hipótesis hasta ejecutar un piloto con tareas comparables y medir tiempo de diagnóstico, falsa alerta y omisiones.")
add_heading(doc,'5. Recomendaciones',1)
add_bullets(doc,[
'Validar las etiquetas con dos revisores de dominio y registrar desacuerdos.',
'Expandir la muestra con mensajes sintéticos de plantillas nuevas y luego con logs reales anonimizados.',
'Comparar Random Forest con regresión logística y Naive Bayes sobre prueba temporal o por grupo, no solo un split aleatorio.',
'Medir la latencia del pipeline completo, incluidos lectura, sanitización, vectorización y emisión de alertas.'
])
doc.add_page_break()
add_heading(doc,'6. Muestra de 25 predicciones para auditoría',1)
add_p(doc,'La siguiente tabla resume una selección estratificada de la partición de prueba. “Coincide” compara la predicción con la etiqueta del CSV; la revisión humana de los casos debe confirmar si la etiqueta de referencia describe correctamente el contenido técnico. La etiqueta sintética no constituye por sí misma un ground truth independiente.')
aud_rows=[]
for _,r in audit.iterrows():
    msg=str(r['mensaje'])
    msg=(msg[:42]+'…') if len(msg)>43 else msg
    aud_rows.append([f"#{int(r['source_row']):03d}",r['severity_real'],r['severity_predicha'],'Sí' if str(r['coincide']).lower()=='true' else 'No',f"{float(r['confianza_prediccion']):.2f}",msg])
add_doc_table(doc,['ID','Real','Predicción','Coincide','Conf.','Fragmento de log sanitizado'],aud_rows,[.55,.8,.9,.65,.55,3.1],7.0)
save_doc(doc,'05_EVALUACION_RESULTADOS.docx')

# --------------------- 06 ethics/business ---------------------
doc=base_doc('06','Análisis ético, gobernanza y viabilidad de negocio','Privacidad, sesgo, supervisión humana, riesgos y escenario económico')
add_callout(doc,'Conclusión ejecutiva','El valor potencial de SpringLog-IA es acelerar la priorización de logs sin ceder el control operativo al modelo. La viabilidad técnica requiere validación en logs reales; las cifras económicas de este informe son un escenario ilustrativo basado en supuestos explícitos, no beneficios observados.')
add_heading(doc,'1. Privacidad y gobernanza de datos',1)
add_p(doc,'Los logs de backend pueden contener JWT/Bearer, tokens, contraseñas, secretos, correos, identificadores de usuario y datos de tarjeta. El script del proyecto incluye reglas Regex deterministas para enmascarar patrones comunes antes de vectorizar los mensajes; estas reglas reducen riesgo, pero no garantizan detección de todos los datos sensibles.')
add_doc_table(doc,['Riesgo','Control en el prototipo','Limitación residual'],[
('JWT, bearer o tokens','Reemplazo por marcadores REDACTED.','Tokens con formatos no previstos pueden escapar.'),
('Contraseñas y API keys','Reglas para campos clave-valor comunes.','Revisar variantes y formatos multilinea en una implementación real.'),
('Correos, IDs y PAN','Sustitución de email, userId/customerId y card/PAN.','Las regex no clasifican por contexto y pueden no cubrir todas las variantes.'),
('Envío a servicios externos','Recomendación de sanitizar localmente antes de invocar APIs.','La exploración generativa no es una integración desplegada.'),
('Retención y acceso','Propuesta de control de acceso, minimización y retención limitada.','Debe concretarse en políticas y controles de la organización.')],[1.4,2.5,2.75],7.9)
add_p(doc,'El diseño toma como referencias la Ley 1581 de 2012 y principios de seguridad y minimización de datos. Este informe no certifica cumplimiento legal; la aplicación concreta debe revisarse con responsables de privacidad y seguridad según el contexto del tratamiento.')
add_heading(doc,'2. Sesgos, error y supervisión humana',1)
add_p(doc,'El dataset tiene 60 % INFO y solo 5 % CRITICAL. El uso de class_weight=balanced y el reporte de F1 macro ayudan a observar el rendimiento por clase, pero no eliminan el sesgo de una muestra sintética. Además, las clases están asociadas a patrones técnicos que pueden ser demasiado fáciles de aprender.')
add_bullets(doc,[
'Human-in-the-loop: las etiquetas son recomendaciones; los incidentes críticos y las predicciones de baja confianza se revisan por un ingeniero.',
'Explicabilidad: conservar el log sanitizado, la clase predicha, la confianza y la versión del modelo; no interpretar la clasificación como causa raíz.',
'Gobernanza: asignar responsable de modelo y responsable operativo, versionar datos/código y definir procedimiento de rollback.',
'Monitoreo: medir falsos negativos de severidad, tasa de WARN degradados a INFO, latencia y cambios en la distribución de logs.'
])
add_heading(doc,'3. Matriz de riesgos',1)
add_doc_table(doc,['Riesgo','Prob.','Impacto','Mitigación / responsable'],[
('WARN degradado a INFO','Alta','Medio-alto','Ajustar features y umbrales; revisar por SRE.'),
('Falso negativo CRITICAL con logs reales','Media','Alto','No confiar en el modelo como único alerta; mantener reglas de guardia.'),
('Datos sensibles no detectados por regex','Media','Alto','Ampliar pruebas, validar con seguridad, bloquear salida en duda.'),
('Degradación por cambio de plantilla','Alta','Medio','Monitorizar deriva y reentrenar con conjunto etiquetado.'),
('ROI sobreestimado','Alta','Medio','Medir piloto y usar costos/volúmenes reales; finanzas/operaciones.'),
('Modelo no reproducible','Baja-media','Medio','Versionar dataset, dependencias, semilla, métricas y artefacto.')],[2.0,.75,.9,3.0],7.8)
add_heading(doc,'4. Escenario económico ilustrativo',1)
add_p(doc,'Para explorar la magnitud del beneficio potencial se reutilizan supuestos de trabajo del proyecto: cinco desarrolladores, 40 incidentes al mes por desarrollador, 35 minutos recuperados por incidente y un costo de oportunidad de USD 35 por hora. El escenario no se ha validado con mediciones reales de SpringLog-IA.')
roi_rows=[
('Minutos recuperados al mes','5 × 40 × 35 = 7.000 min'),
('Horas recuperadas al mes','7.000 / 60 = 116,67 h'),
('Ahorro bruto mensual','116,67 h × USD 35 = USD 4.083,33'),
('Costo operativo mensual supuesto','USD 50,00'),
('Ahorro neto mensual estimado','USD 4.033,33'),
('Ahorro neto anual antes de inversión','USD 4.033,33 × 12 = USD 48.400,00'),
('Inversión inicial supuesta','USD 1.500,00'),
('Beneficio primer año después de inversión','USD 48.400 − USD 1.500 = USD 46.900'),
('ROI simple del primer año','USD 46.900 / USD 1.500 ≈ 3.127 %'),
('Payback aproximado','USD 1.500 / USD 4.033,33 × 30 ≈ 11,2 días')]
add_doc_table(doc,['Concepto','Cálculo'],roi_rows,[2.55,3.95],8.2)
add_callout(doc,'Condición de uso','Estas cifras no deben presentarse como resultados medidos. El beneficio depende, entre otras variables, del volumen real de incidentes, del tiempo realmente ahorrado, de falsos positivos/negativos y del costo operativo efectivo.')
add_heading(doc,'5. Sostenibilidad y viabilidad técnica',1)
add_p(doc,'Un clasificador scikit-learn ejecutado localmente puede evitar enviar cada log a un LLM externo, pero no se ha medido energía ni huella de carbono en este proyecto. No se afirma un porcentaje de reducción de CO₂e. Una futura comparación Green AI debe medir ambos enfoques bajo las mismas tareas, hardware, carga, calidad de salida y ventana de tiempo.')
add_heading(doc,'6. Plan de implementación',1)
add_bullets(doc,[
'Fase 1 — Validación offline: ampliar dataset y etiquetas; medir métricas por clase y por plantilla.',
'Fase 2 — Piloto sombreado: procesar logs ya sanitizados sin afectar alertas de producción y comparar con decisiones de SRE.',
'Fase 3 — Observabilidad: exponer métricas del clasificador, registrar versión/confianza y definir umbrales revisados por operación.',
'Fase 4 — Decisión de adopción: evaluar MTTR real, costo, privacidad, tasa de fallos y beneficio neto antes de activar alertas asistidas.'
])
save_doc(doc,'06_ANALISIS_ETICO_NEGOCIO.docx')

# --------------------- Excel workbook ---------------------
wb=Workbook(); ws=wb.active; ws.title='Resumen'
ws.sheet_view.showGridLines=False
thin=Side(style='thin',color='D5DEE8')
def apply_title(ws,title,subtitle,last_col=7):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=last_col)
    c=ws.cell(1,1,title); c.font=Font(name='Aptos Display',size=18,bold=True,color=WHITE); c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=32
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=last_col)
    c=ws.cell(2,1,subtitle); c.font=Font(name='Aptos',size=10,color=MUTED,italic=True); c.alignment=Alignment(wrap_text=True,vertical='center'); ws.row_dimensions[2].height=28

def header_row(ws,row,headers):
    for i,h in enumerate(headers,1):
        c=ws.cell(row,i,h); c.font=Font(name='Aptos',size=9,bold=True,color=WHITE); c.fill=PatternFill('solid',fgColor=TEAL); c.alignment=Alignment(wrap_text=True,vertical='center'); c.border=Border(bottom=thin)
    ws.row_dimensions[row].height=25

def style_body(ws,min_row,max_row,min_col,max_col):
    for row in ws.iter_rows(min_row=min_row,max_row=max_row,min_col=min_col,max_col=max_col):
        for c in row:
            c.font=Font(name='Aptos',size=9,color=TEXT)
            c.alignment=Alignment(vertical='top',wrap_text=True)
            c.border=Border(bottom=Side(style='hair',color='E1E8EF'))
            if c.row%2==0: c.fill=PatternFill('solid',fgColor='F6F8FB')

def add_excel_table(ws,ref,name):
    tab=Table(displayName=name,ref=ref); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False); ws.add_table(tab)

apply_title(ws,'SpringLog-IA | Análisis exploratorio de logs','Fuente: 500 registros sintéticos de payment-service. Las frecuencias son descriptivas, no incidentes observados en producción.',8)
ws['A4']='Indicador'; ws['B4']='Valor'; header_row(ws,4,['Indicador','Valor'])
summary_rows=[('Registros',len(clean)),('Columnas originales',8),('Columnas limpias',len(clean.columns)),('Duplicados exactos',eda['exact_duplicates']),('Marcadores None normalizados',eda['source_exception_none_sentinel']),('Inicio',eda['timestamp_start']),('Fin',eda['timestamp_end']),('Modelo de referencia','Random Forest + TF-IDF'),('Accuracy de prueba',metrics['accuracy']),('F1 macro de prueba',metrics['f1_macro']),('P95 inferencia local (ms)',metrics['latency']['p95_ms'])]
for r,(k,v) in enumerate(summary_rows,5): ws.cell(r,1,k); ws.cell(r,2,v)
style_body(ws,5,4+len(summary_rows),1,2); ws.column_dimensions['A'].width=34; ws.column_dimensions['B'].width=44
ws['B13'].number_format='0.0%'; ws['B14'].number_format='0.000'; ws['B15'].number_format='0.0'
# severity distribution plus chart
sev=wb.create_sheet('Severidad'); sev.sheet_view.showGridLines=False; apply_title(sev,'Distribución de severidad','Distribución definida al generar la muestra; revisar desbalance de clases al interpretar las métricas.',4)
header_row(sev,4,['Severidad','Registros','Porcentaje','Interpretación'])
sev_notes={'INFO':'Tráfico rutinario','WARN':'Advertencia/latencia','ERROR':'Excepción/error de aplicación','CRITICAL':'Evento crítico sintético'}
for i,s in enumerate(['INFO','WARN','ERROR','CRITICAL'],5):
    sev.cell(i,1,s); sev.cell(i,2,int(eda['severity_counts'][s])); sev.cell(i,3,float(eda['severity_percent'][s])/100); sev.cell(i,4,sev_notes[s]); sev.cell(i,3).number_format='0.0%'
style_body(sev,5,8,1,4); [setattr(sev.column_dimensions[c],'width',w) for c,w in [('A',18),('B',14),('C',14),('D',36)]]
chart=BarChart(); chart.type='col'; chart.style=10; chart.title='Registros por severidad'; chart.y_axis.title='Registros'; chart.x_axis.title='Severidad'; chart.height=7; chart.width=13
chart.add_data(Reference(sev,min_col=2,min_row=4,max_row=8),titles_from_data=True); chart.set_categories(Reference(sev,min_col=1,min_row=5,max_row=8)); chart.legend=None; sev.add_chart(chart,'F4')
# Instance distribution
inst=wb.create_sheet('Instancias'); inst.sheet_view.showGridLines=False; apply_title(inst,'Severidad por instancia','Conteos observados para las tres instancias sintéticas.',7)
header_row(inst,4,['Instancia','INFO','WARN','ERROR','CRITICAL','Total','Porcentaje'])
for i,node in enumerate(['instance-prod-east-01','instance-prod-east-02','instance-prod-west-01'],5):
    sub=clean[clean.instance_id==node]; vals=[int((sub.severity==s).sum()) for s in ['INFO','WARN','ERROR','CRITICAL']]
    inst.cell(i,1,node)
    for j,v in enumerate(vals,2): inst.cell(i,j,v)
    inst.cell(i,6,f'=SUM(B{i}:E{i})'); inst.cell(i,7,f'=F{i}/SUM($F$5:$F$7)'); inst.cell(i,7).number_format='0.0%'
style_body(inst,5,7,1,7)
for col,width in zip('ABCDEFG',[32,12,12,12,12,12,14]): inst.column_dimensions[col].width=width
chart=BarChart(); chart.type='bar'; chart.grouping='stacked'; chart.overlap=100; chart.style=12; chart.title='Composición de severidad por instancia'; chart.x_axis.title='Registros'; chart.y_axis.title='Instancia'; chart.height=7; chart.width=15
chart.add_data(Reference(inst,min_col=2,max_col=5,min_row=4,max_row=7),titles_from_data=True); chart.set_categories(Reference(inst,min_col=1,min_row=5,max_row=7)); chart.legend.position='b'; inst.add_chart(chart,'I4')
# Exceptions
ex=wb.create_sheet('Excepciones'); ex.sheet_view.showGridLines=False; apply_title(ex,'Tipos de excepción frecuentes','Se cuentan valores de exception_type distintos de NO_APLICA.',3)
header_row(ex,4,['Tipo de excepción','Registros','Nota'])
for i,(name,count) in enumerate(eda['top_exceptions'].items(),5):
    ex.cell(i,1,name); ex.cell(i,2,count); ex.cell(i,3,'Conteo del dataset sintético')
style_body(ex,5,4+len(eda['top_exceptions']),1,3)
ex.column_dimensions['A'].width=42; ex.column_dimensions['B'].width=14; ex.column_dimensions['C'].width=36
bar=BarChart(); bar.type='bar'; bar.style=10; bar.title='Excepciones más frecuentes'; bar.x_axis.title='Registros'; bar.height=8; bar.width=15
bar.add_data(Reference(ex,min_col=2,min_row=4,max_row=4+len(eda['top_exceptions'])),titles_from_data=True); bar.set_categories(Reference(ex,min_col=1,min_row=5,max_row=4+len(eda['top_exceptions']))); bar.legend=None; ex.add_chart(bar,'E4')
# Data quality
q=wb.create_sheet('Calidad'); q.sheet_view.showGridLines=False; apply_title(q,'Calidad y transformación','La limpieza distingue nulos reales de marcadores textuales y aplica sanitización.',4)
header_row(q,4,['Prueba','Resultado','Acción','Estado'])
quality=[('Registros',500,'Comprobar número esperado','OK'),('Columnas obligatorias presentes',8,'Validar esquema','OK'),('Duplicados exactos',eda['exact_duplicates'],'Contar registros repetidos','OK' if eda['exact_duplicates']==0 else 'Revisar'),('Marcadores literales None en fuente',eda['source_exception_none_sentinel'],'Convertir a NO_APLICA','Normalizado'),('Nulos tras limpieza',0,'Validar valores vacíos','OK'),('JWT/tokens/secretos','Sanitización basada en regex','Reemplazar con REDACTED','Implementado en pipeline'),('IDs/correos/PAN','Sanitización basada en regex','Reemplazar por marcadores','Implementado en pipeline'),('Menciones de severidad en texto','Enmascaradas en características','Evitar atajo directo de etiqueta','Implementado en pipeline')]
for i,row in enumerate(quality,5):
    for j,v in enumerate(row,1): q.cell(i,j,v)
style_body(q,5,4+len(quality),1,4)
for col,width in zip('ABCD',[35,34,42,24]): q.column_dimensions[col].width=width
# Model comparison
mc=wb.create_sheet('Modelos'); mc.sheet_view.showGridLines=False; apply_title(mc,'Comparación de modelos','Partición estratificada 80/20, random_state=42; evaluación sobre datos sintéticos.',6)
header_row(mc,4,['Modelo','Accuracy','Precision macro','Recall macro','F1 macro','F1 ponderado'])
for i,(_,r) in enumerate(model_cmp.iterrows(),5):
    vals=[r['Modelo'],r['Accuracy'],r['Precision_macro'],r['Recall_macro'],r['F1_macro'],r['F1_weighted']]
    for j,v in enumerate(vals,1): mc.cell(i,j,v)
    for j in range(2,7): mc.cell(i,j).number_format='0.0%'
style_body(mc,5,4+len(model_cmp),1,6)
for col,width in zip('ABCDEF',[34,14,17,15,12,16]): mc.column_dimensions[col].width=width
mc['A11']='Nota'; mc['B11']='Regresión logística y Naive Bayes al 100 % en esta muestra requieren una validación de generalización con plantillas nuevas.'; mc.merge_cells('B11:F12'); mc['B11'].alignment=Alignment(wrap_text=True,vertical='top'); mc['B11'].fill=PatternFill('solid',fgColor='FFF5DB'); mc['B11'].font=Font(name='Aptos',size=9,color=TEXT)
# Clean data sheet
wsd=wb.create_sheet('Datos_limpios'); wsd.sheet_view.showGridLines=False
for j,col in enumerate(clean.columns,1): wsd.cell(1,j,col)
header_row(wsd,1,list(clean.columns))
for i,row in enumerate(clean.itertuples(index=False,name=None),2):
    for j,val in enumerate(row,1): wsd.cell(i,j,val)
style_body(wsd,2,len(clean)+1,1,len(clean.columns))
for col,width in {'A':22,'B':18,'C':28,'D':14,'E':24,'F':46,'G':30,'H':70,'I':14,'J':10,'K':16}.items(): wsd.column_dimensions[col].width=width
wsd.freeze_panes='A2'; wsd.auto_filter.ref=f'A1:{chr(64+len(clean.columns))}{len(clean)+1}'
add_excel_table(wsd,f'A1:K{len(clean)+1}','DatosLimpiosSpringLog')
for row in wsd.iter_rows(min_row=2,max_row=len(clean)+1,min_col=1,max_col=11):
    row[0].number_format='yyyy-mm-dd hh:mm:ss'
# workbook general formatting/freeze panes
for sh in wb.worksheets:
    sh.sheet_properties.pageSetUpPr.fitToPage=True
    sh.page_setup.fitToWidth=1; sh.page_setup.fitToHeight=0; sh.page_margins.left=.25; sh.page_margins.right=.25
    if sh.title not in ('Datos_limpios',): sh.freeze_panes='A5'
    sh.sheet_properties.tabColor=TEAL if sh.title in ('Resumen','Datos_limpios') else NAVY
wb.save(ROOT/'03_ANALISIS_DATOS.xlsx')

# --------------------- PPTX ---------------------
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs=Presentation(); prs.slide_width=PInches(13.333); prs.slide_height=PInches(7.5)
P_NAVY=PRGB(22,50,79); P_TEAL=PRGB(15,139,141); P_MUTED=PRGB(91,101,115); P_TEXT=PRGB(36,52,71); P_LIGHT=PRGB(234,242,248); P_WHITE=PRGB(255,255,255); P_AMBER=PRGB(224,161,38); P_RED=PRGB(217,93,57)

def set_bg(slide,color=P_WHITE):
    fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=color

def textbox(slide,x,y,w,h,text,size=16,color=P_TEXT,bold=False,font='Aptos',align=PP_ALIGN.LEFT,margin=0.05):
    shape=slide.shapes.add_textbox(PInches(x),PInches(y),PInches(w),PInches(h)); tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=PInches(margin); tf.margin_right=PInches(margin); tf.margin_top=PInches(margin); tf.margin_bottom=PInches(margin)
    p=tf.paragraphs[0]; p.text=str(text); p.alignment=align
    p.font.name=font; p.font.size=PPt(size); p.font.bold=bold; p.font.color.rgb=color
    return shape

def base_slide(title,number,section='PROYECTO FINAL INTEGRADO'):
    s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
    # thin brand bar and headers
    bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,PInches(.11)); bar.fill.solid(); bar.fill.fore_color.rgb=P_TEAL; bar.line.fill.background()
    textbox(s,.55,.24,8,.25,'DIPLOMADO EN INTELIGENCIA ARTIFICIAL  |  '+section,9,P_TEAL,True)
    textbox(s,.55,.63,12.2,.53,title,25,P_NAVY,True,'Aptos Display')
    line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,PInches(.55),PInches(1.28),PInches(12.2),PInches(.025)); line.fill.solid(); line.fill.fore_color.rgb=P_TEAL; line.line.fill.background()
    textbox(s,.55,7.15,10,.18,'SPRINGLOG-IA  |  Observabilidad asistida de logs Spring Boot',8,P_MUTED)
    textbox(s,12.15,7.13,.6,.2,f'{number:02d}',8,P_MUTED,True,align=PP_ALIGN.RIGHT)
    return s

def bullet_box(slide,x,y,w,h,items,font_size=15,color=P_TEXT,spacing=7):
    shape=slide.shapes.add_textbox(PInches(x),PInches(y),PInches(w),PInches(h)); tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=PInches(.12); tf.margin_right=PInches(.12); tf.margin_top=PInches(.1); tf.margin_bottom=PInches(.1)
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=item; p.level=0; p.font.name='Aptos'; p.font.size=PPt(font_size); p.font.color.rgb=color; p.space_after=PPt(spacing); p.text='•  '+item
    return shape

def metric_card(slide,x,y,w,h,value,label,accent=P_TEAL):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PInches(x),PInches(y),PInches(w),PInches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=P_LIGHT; sh.line.color.rgb=accent; sh.line.width=PPt(1.4)
    textbox(slide,x+.12,y+.12,w-.24,.48,value,24,accent,True,'Aptos Display',PP_ALIGN.CENTER)
    textbox(slide,x+.12,y+.68,w-.24,h-.78,label,11,P_TEXT,False,align=PP_ALIGN.CENTER)

def add_ppt_table(slide,x,y,w,h,headers,rows,widths=None,font_size=11):
    shape=slide.shapes.add_table(len(rows)+1,len(headers),PInches(x),PInches(y),PInches(w),PInches(h)); tb=shape.table
    if widths:
        for i,wi in enumerate(widths): tb.columns[i].width=PInches(wi)
    for ri,row in enumerate([headers]+rows):
        for ci,val in enumerate(row):
            cell=tb.cell(ri,ci); cell.text=str(val); cell.margin_left=PInches(.06); cell.margin_right=PInches(.06); cell.margin_top=PInches(.04); cell.margin_bottom=PInches(.04)
            cell.fill.solid(); cell.fill.fore_color.rgb=P_NAVY if ri==0 else (P_LIGHT if ri%2 else P_WHITE)
            for p in cell.text_frame.paragraphs:
                p.font.name='Aptos'; p.font.size=PPt(font_size); p.font.bold=(ri==0); p.font.color.rgb=P_WHITE if ri==0 else P_TEXT
    return shape

# Slide 1
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s,P_NAVY)
bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,PInches(.15)); bar.fill.solid(); bar.fill.fore_color.rgb=P_TEAL; bar.line.fill.background()
textbox(s,.8,.85,10,.35,'DIPLOMADO EN INTELIGENCIA ARTIFICIAL  |  OPCIÓN DE GRADO',12,P_TEAL,True)
textbox(s,.8,1.75,11.5,1.15,'SpringLog-IA',38,P_WHITE,True,'Aptos Display')
textbox(s,.8,2.95,10.7,1.0,'Clasificación asistida de severidad de logs backend en Java / Spring Boot',21,P_LIGHT,False)
line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,PInches(.8),PInches(4.28),PInches(2.2),PInches(.07)); line.fill.solid(); line.fill.fore_color.rgb=P_TEAL; line.line.fill.background()
textbox(s,.8,4.72,8,.5,'David Felipe Rios Echeverri',16,P_WHITE,True)
textbox(s,.8,5.2,8,.4,'Desarrollo Backend / Ingeniería de Software  •  Octubre de 2026',12,P_LIGHT)
textbox(s,.8,6.85,10,.25,'Proyecto integrado con datos sintéticos, validación reproducible y supervisión humana.',10,P_LIGHT)
# 2 Executive summary
s=base_slide('Resumen ejecutivo: qué se construyó y qué demuestra',2)
metric_card(s,.7,1.75,2.8,1.35,'93 %','Accuracy en 100 registros de prueba')
metric_card(s,3.75,1.75,2.8,1.35,'0,933','F1 macro del Random Forest')
metric_card(s,6.8,1.75,2.8,1.35,'5/5','CRITICAL detectados en test',P_RED)
metric_card(s,9.85,1.75,2.8,1.35,'7','WARN clasificados como INFO',P_AMBER)
bullet_box(s,.8,3.55,11.8,2.6,[
'Pipeline local: sanitización basada en reglas → TF-IDF → Random Forest.',
'La evaluación se hizo con un split estratificado 80/20 y semilla fija.',
'El dataset sintético es útil para demostrar el flujo, pero no acredita calidad de producción.',
'El siguiente hito es validar en plantillas nuevas y con revisión de especialistas.'
],15)
# 3 Problem
s=base_slide('Problema: encontrar la señal entre miles de eventos',3)
textbox(s,.85,1.65,5.4,.45,'SITUACIÓN ACTUAL',13,P_TEAL,True)
bullet_box(s,.85,2.1,5.35,3.7,['Logs estructurados y mensajes no estructurados se mezclan.','El análisis manual de stack traces ralentiza la priorización.','WARN puede competir por atención con eventos rutinarios.','Los filtros por palabras clave no siempre captan el contexto.'],15)
textbox(s,7.0,1.65,5.2,.45,'NECESIDAD',13,P_TEAL,True)
bullet_box(s,7.0,2.1,5.25,3.7,['Clasificar severidad como ayuda al analista.','Sanitizar información sensible antes de procesarla.','Priorizar señales críticas sin eliminar la supervisión humana.','Medir resultados con un conjunto de prueba reproducible.'],15)
# 4 Evidence vs hypotheses
s=base_slide('Separar el punto de partida de los resultados medidos',4)
add_ppt_table(s,.8,1.7,11.8,3.9,['Afirmación','Tipo de evidencia','Estado en este proyecto'],[
('MTTR inicial de 45 min','Hipótesis inicial del proyecto','No medido en producción'),
('Reducción del MTTR a 10 min','Escenario de beneficio','Pendiente de piloto'),
('Accuracy 93 %','Prueba sintética reservada n=100','Medido en el split actual'),
('Recall CRITICAL 100 %','5 de 5 casos en test','Medido; soporte pequeño'),
(f"P95 inferencia {metrics['latency']['p95_ms']:.1f} ms",'Medición local sobre 30 registros','Dependiente del entorno')],widths=[3.2,4.0,4.6],font_size=12)
textbox(s,.9,5.95,11.5,.5,'Regla de comunicación: no presentar estimaciones económicas o hipótesis operativas como resultados experimentales.',14,P_NAVY,True)
# 5 data profile
s=base_slide('Datos: una muestra sintética de cinco días',5)
metric_card(s,.8,1.65,2.5,1.15,'500','registros; 8 campos originales')
metric_card(s,3.5,1.65,2.5,1.15,'3','instancias del microservicio')
metric_card(s,6.2,1.65,2.5,1.15,'0','duplicados exactos')
metric_card(s,8.9,1.65,3.3,1.15,'389','marcadores None normalizados')
s.shapes.add_picture(str(CHART/'eda_severidad.png'),PInches(.9),PInches(3.0),width=PInches(5.75))
bullet_box(s,7.0,3.1,5.0,2.8,['INFO: 60 %','WARN: 20 %','ERROR: 15 %','CRITICAL: 5 %','Proporciones fijadas al generar la muestra.'],15)
# 6 node and exception analysis
s=base_slide('EDA: variación entre instancias y tipos de excepción',6)
s.shapes.add_picture(str(CHART/'eda_instancias.png'),PInches(.65),PInches(1.6),width=PInches(6.1))
s.shapes.add_picture(str(CHART/'eda_excepciones.png'),PInches(6.9),PInches(1.6),width=PInches(5.8))
textbox(s,.85,6.1,11.3,.4,'instance-prod-west-01 tiene el mayor volumen de muestra (193) y 13 registros CRITICAL; estos conteos no reflejan una infraestructura real.',11,P_MUTED)
# 7 architecture
s=base_slide('Propuesta técnica: pipeline local de clasificación',7)
steps=[('1','Ingesta','CSV y contexto JVM'),('2','Sanitización','Regex: secretos, IDs, correos'),('3','TF-IDF','Unigramas + bigramas'),('4','Clasificador','Random Forest · 100 árboles'),('5','Revisión','Etiqueta y confianza para analista')]
for i,(n,t,d) in enumerate(steps):
    x=.65+i*2.55
    box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PInches(x),PInches(2.25),PInches(2.2),PInches(2.0)); box.fill.solid(); box.fill.fore_color.rgb=P_LIGHT; box.line.color.rgb=P_TEAL
    textbox(s,x+.12,2.45,1.95,.35,n,22,P_TEAL,True,align=PP_ALIGN.CENTER)
    textbox(s,x+.12,2.88,1.95,.35,t,15,P_NAVY,True,align=PP_ALIGN.CENTER)
    textbox(s,x+.12,3.35,1.95,.72,d,10,P_TEXT,False,align=PP_ALIGN.CENTER)
    if i<4:
        tx=s.shapes.add_shape(MSO_SHAPE.CHEVRON,PInches(x+2.23),PInches(2.92),PInches(.28),PInches(.45)); tx.fill.solid(); tx.fill.fore_color.rgb=P_TEAL; tx.line.fill.background()
textbox(s,.9,5.0,11.4,.8,'El prototipo clasifica severidad. No ejecuta remediaciones, no establece causa raíz y todavía no se conecta en tiempo real a Logback/Log4j2.',15,P_NAVY,True,align=PP_ALIGN.CENTER)
# 8 GenAI
s=base_slide('IA generativa: utilidad exploratoria y límites',8)
add_ppt_table(s,.7,1.65,6.0,4.4,['Flujo','Uso esperado'],[
('Extracción JSON','Normalizar campos y resumir una línea.'),('Diagnóstico asistido','Separar hechos, hipótesis y comprobaciones.'),('Sanitización','Proponer/validar enmascaramiento; no único control.'),('Resumen operativo','Contar severidades y describir patrones con cautela.')],widths=[2.1,3.9],font_size=12)
bullet_box(s,7.2,1.9,5.0,3.6,['Una salida plausible no demuestra causa raíz.','La sanitización debe ser determinista antes de una API externa.','El costo y la latencia dependen del volumen y del modelo.','Versionar prompts y guardar casos de prueba.'],14)
# 9 Model setup
s=base_slide('Modelo: entrenamiento y evaluación reproducibles',9)
metric_card(s,.85,1.7,2.5,1.2,'400','registros entrenamiento')
metric_card(s,3.55,1.7,2.5,1.2,'100','registros prueba')
metric_card(s,6.25,1.7,2.5,1.2,'1.000','máx. características TF-IDF')
metric_card(s,8.95,1.7,3.2,1.2,'42','semilla de partición')
bullet_box(s,1.0,3.45,11.2,2.4,['TF-IDF, n-gramas 1–2, sublinear_tf=True.','Random Forest: 100 árboles y class_weight=balanced.','Se compararon también Regresión logística, Naive Bayes y Árbol de decisión.','El pipeline y sus métricas se guardan para reproducir la evaluación.'],15)
# 10 Model comparison
s=base_slide('Comparación: los modelos simples también puntúan alto',10)
s.shapes.add_picture(str(CHART/'modelo_comparacion.png'),PInches(.65),PInches(1.55),width=PInches(7.1))
add_ppt_table(s,8.0,1.9,4.5,3.7,['Modelo','Accuracy','F1 macro'],[
('Random Forest','93 %','0,933'),('Regresión logística','100 %','1,000'),('Naive Bayes','100 %','1,000'),('Árbol de decisión','98 %','0,983')],widths=[2.35,1.05,1.1],font_size=11)
textbox(s,8.05,5.85,4.35,.7,'Las puntuaciones perfectas sugieren patrones sintéticos demasiado separables; requieren validación con plantillas nuevas.',10,P_RED,True)
# 11 results/confusion
s=base_slide('Resultado del Random Forest y principal caso borde',11)
s.shapes.add_picture(str(CHART/'modelo_matriz_confusion.png'),PInches(.75),PInches(1.55),width=PInches(5.6))
metric_card(s,6.85,1.8,2.2,1.2,'5/5','CRITICAL correctos',P_RED)
metric_card(s,9.25,1.8,2.5,1.2,'15/15','ERROR correctos',P_RED)
textbox(s,6.85,3.45,5.0,.55,'WARN: 13/20 detectados (65 % recall)',17,P_NAVY,True)
bullet_box(s,6.85,4.05,5.2,1.9,['7 WARN se clasificaron como INFO.','Explorar durationMs, estado HTTP y reintentos como variables estructuradas.','El test tiene solo 5 casos CRITICAL.'],12)
# 12 security
s=base_slide('Gobernanza: privacidad, sesgo y supervisión humana',12)
add_ppt_table(s,.75,1.65,11.85,3.75,['Control','Qué se hace','Qué falta'],[
('Privacidad','Regex para secretos/tokens, correos, IDs y PAN.','Pruebas con variantes y revisión de seguridad.'),
('Supervisión humana','Modelo recomendado, no autónomo.','Procedimiento formal de revisión y rollback.'),
('Sesgo de clase','Class_weight=balanced y F1 macro.','Más ejemplos CRITICAL y validación externa.'),
('Explicabilidad','Conservar texto sanitizado y clase predicha.','Documentar versión, características y razones de alerta.')],widths=[2.05,4.35,5.45],font_size=11)
textbox(s,.9,5.8,11.1,.5,'No se afirma cumplimiento legal certificado ni detección total de secretos mediante Regex.',13,P_RED,True,align=PP_ALIGN.CENTER)
# 13 business case
s=base_slide('Viabilidad: escenario económico basado en supuestos',13)
metric_card(s,.8,1.7,2.7,1.3,'USD 4.033','ahorro mensual neto estimado')
metric_card(s,3.75,1.7,2.7,1.3,'USD 46.900','beneficio año 1 tras inversión')
metric_card(s,6.7,1.7,2.7,1.3,'3.127 %','ROI simple del primer año')
metric_card(s,9.65,1.7,2.7,1.3,'11,2 días','payback teórico')
bullet_box(s,1.0,3.65,11.2,2.5,['Supuestos: 5 desarrolladores, 40 incidentes/mes/persona, 35 min recuperados por incidente y USD 35/h.','Costo operativo supuesto: USD 50/mes; inversión inicial: USD 1.500.','MTTR y ahorro de tiempo no se han medido; validar con piloto antes de usar estas cifras como caso de negocio.'],14)
# 14 roadmap
s=base_slide('Hoja de ruta: de demostración a piloto controlado',14)
road=[('Fase 1','Validación offline','Etiquetas revisadas, plantillas nuevas, métricas por clase.'),('Fase 2','Modo sombra','Comparar recomendaciones con el flujo de SRE sin activar alertas.'),('Fase 3','Integración','Métricas, trazabilidad de versión, umbrales y fallback.'),('Fase 4','Decisión','Medir MTTR, falsos negativos, costo y seguridad para decidir adopción.')]
for i,(n,t,d) in enumerate(road):
    y=1.62+i*1.22
    circ=s.shapes.add_shape(MSO_SHAPE.OVAL,PInches(.95),PInches(y+.05),PInches(.55),PInches(.55)); circ.fill.solid(); circ.fill.fore_color.rgb=P_TEAL; circ.line.fill.background()
    textbox(s,1.0,y+.13,.45,.25,str(i+1),12,P_WHITE,True,align=PP_ALIGN.CENTER)
    textbox(s,1.8,y,2.2,.4,n+' — '+t,15,P_NAVY,True)
    textbox(s,4.3,y,7.8,.55,d,13,P_TEXT)
# 15 conclusion
s=base_slide('Conclusiones: prototipo útil, validación pendiente',15)
bullet_box(s,.95,1.72,11.4,3.7,[
'SpringLog-IA demuestra un flujo reproducible de clasificación textual de severidad para payment-service sintético.',
'Random Forest consigue 93 % de accuracy y 0,933 de F1 macro en la partición usada; la clase WARN es el principal punto de mejora.',
'Los resultados del dataset no bastan para asegurar generalización ni reducción del MTTR real.',
'El siguiente paso clave es validar con plantillas nuevas, revisión de especialistas y piloto de modo sombra.'
],16)
textbox(s,1.0,5.9,11.3,.5,'Preguntas',23,P_TEAL,True,align=PP_ALIGN.CENTER)
prs.save(ROOT/'07_PRESENTACION_EJECUTIVA.pptx')

# --------------------- Supporting docs / README ---------------------
# Short narration script to support the optional 3-5 min executive summary video.
doc=base_doc('07-A','Guion para resumen ejecutivo','Guion opcional de narración para video de 3–5 minutos')
add_callout(doc,'Duración objetivo','Aproximadamente 4 minutos a ritmo de exposición conversacional. El guion no afirma que el video haya sido grabado.')
segments=[
('0:00–0:35 | Problema','SpringLog-IA nace de un problema de operación backend: al investigar incidentes, los equipos necesitan distinguir el tráfico rutinario de advertencias y fallas relevantes dentro de logs no estructurados. Los filtros por cadenas ayudan, pero el contexto técnico puede quedar disperso entre mensajes e instancias.'),
('0:35–1:05 | Datos','La demostración usa 500 registros sintéticos de payment-service recogidos en una ventana de cinco días y distribuidos entre tres instancias. Las clases están desbalanceadas: 60 % INFO y solo 5 % CRITICAL. Esto permite probar el proceso, no representar la frecuencia real de incidentes.'),
('1:05–1:50 | Solución','El pipeline normaliza valores ausentes, sanitiza identificadores y secretos mediante reglas, representa el texto con TF-IDF y aplica Random Forest para sugerir una severidad. La salida está destinada al análisis humano; no ejecuta cambios ni reemplaza la observabilidad existente.'),
('1:50–2:40 | Resultados','En una partición reservada de 100 registros, el Random Forest alcanza 93 % de accuracy y un F1 macro cercano a 0,933. Detecta los cinco ejemplos CRITICAL y los quince ERROR del test, pero deja pasar siete de veinte WARN como INFO. La muestra es pequeña y sintética.'),
('2:40–3:25 | Limitaciones y responsabilidad','Otros clasificadores alcanzaron puntuaciones perfectas sobre la misma muestra, una señal de que los patrones sintéticos son demasiado separables. Por eso no se puede asegurar generalización a producción ni afirmar que se reduce el MTTR. La privacidad depende también de pruebas adicionales de la sanitización y de la revisión humana.'),
('3:25–4:05 | Siguiente paso','La ruta propuesta consiste en validar con plantillas nuevas y etiquetas revisadas, operar primero en modo sombra, medir el tiempo de diagnóstico frente al proceso manual y solo después evaluar integración con las herramientas de observabilidad. El valor del proyecto es un prototipo comprobable y una ruta de validación, no una promesa de producción inmediata.')]
for title,body in segments:
    add_heading(doc,title,2); add_p(doc,body)
save_doc(doc,'07_GUION_RESUMEN_EJECUTIVO.docx')

# README and requirements
readme='''# SpringLog-IA — paquete final de proyecto

## Resumen
Prototipo reproducible de clasificación asistida de severidad para logs sintéticos del microservicio `payment-service`. Incluye análisis exploratorio, sanitización mediante reglas, comparación de clasificadores, modelo serializado, informes y presentación. No es una integración productiva ni una prueba sobre telemetría real.

## Entregables
- `01_FICHA_DEL_PROBLEMA.docx` y `01_FICHA_DEL_PROBLEMA.pdf`: definición del problema, alcance y criterios de éxito.
- `02_IA_GENERATIVA.docx`: prompts, ejemplos de entrada/salida, capacidades, límites y pasos siguientes.
- `03_ANALISIS_DATOS.docx`, `03_ANALISIS_DATOS.pdf` y `03_ANALISIS_DATOS.xlsx`: EDA y dataset preparado; el libro contiene resumen, severidades, instancias, excepciones, calidad, modelos y datos limpios.
- `04_MODELO_O_PROTOTIPO.docx`, `04_MODELO_O_PROTOTIPO.py` y `04_MODELO_O_PROTOTIPO.joblib`: arquitectura, evaluación, script y modelo serializado.
- `05_EVALUACION_RESULTADOS.docx`: métricas reproducibles, matriz de confusión, casos borde y muestra estratificada de 25 predicciones.
- `06_ANALISIS_ETICO_NEGOCIO.docx`: privacidad, sesgos, supervisión, riesgos y escenario económico explícitamente hipotético.
- `07_PRESENTACION_EJECUTIVA.pptx` y `07_PRESENTACION_EJECUTIVA.pdf`: presentación de 15 diapositivas.
- `07_VIDEO_RESUMEN_EJECUTIVO.mp4`: video de aproximadamente 4 minutos y 48 segundos con narración sintética en español, basado en la presentación. Puede reemplazarse por una grabación personal.
- `07_GUION_RESUMEN_EJECUTIVO.docx`: guion de narración para adaptar o volver a grabar.

## Resultados de la última ejecución registrada
- Dataset: 500 registros, 8 columnas de origen y 11 columnas en la versión preparada; 0 duplicados exactos.
- Etiquetas: INFO 300 (60 %), WARN 100 (20 %), ERROR 75 (15 %) y CRITICAL 25 (5 %). La distribución fue fijada al generar datos sintéticos.
- Random Forest: accuracy 93,0 %, F1 macro 0,933 y F1 ponderado 0,925 en una partición estratificada de prueba de 100 filas (`random_state=42`).
- Matriz de confusión: 5/5 CRITICAL y 15/15 ERROR correctos; 13/20 WARN correctos, con 7 WARN clasificados como INFO.
- Latencia local: mediana 21,5 ms y P95 23,6 ms por registro en la ejecución registrada; depende del hardware y no equivale al MTTR.

**Advertencia de validez:** la muestra es pequeña y sintética. Regresión logística y Naive Bayes obtuvieron 100 % en la misma partición, lo que sugiere que algunas plantillas son demasiado separables. Las métricas no demuestran generalización a producción; la severidad crítica tiene solo cinco casos de prueba. La muestra de 25 predicciones es una hoja de trabajo para revisión humana, no una auditoría independiente.

## Requisitos y ejecución
Python 3.10+ recomendado. Instala dependencias con:

```bash
pip install -r requirements.txt
```

Desde la raíz del paquete, ejecuta en orden:

```bash
python 03_ANALISIS_DATOS.py
python 04_MODELO_O_PROTOTIPO.py
```

El primer script valida el CSV de origen, normaliza los 389 marcadores literales `None` como `NO_APLICA`, aplica reglas de sanitización, crea `data/payment_service_logs_500_clean.csv` y `outputs/payment_service_logs_500_clean.csv`, y genera las gráficas y resúmenes del EDA.

El segundo script toma el dataset preparado, construye características desde `logger_class`, `exception_type` y `log_message`, enmascara tokens de severidad dentro del texto, compara cuatro clasificadores, calcula métricas en la partición reservada, genera las gráficas y serializa el pipeline. La evaluación usa el split de prueba; el `.joblib` final se reentrena sobre todos los registros para la demostración posterior.

Para reconstruir informes, libro y presentación desde los últimos resultados guardados, ejecuta opcionalmente:

```bash
python build_deliverables.py
```

## Contrato de inferencia
El `.joblib` recibe una lista de cadenas ya compuestas y sanitizadas con el formato `logger_class + exception_type + log_message`. Consulta las funciones `sanitize` y `build_features` en `04_MODELO_O_PROTOTIPO.py`. No envíes logs sensibles a servicios externos sin sanitización local y revisión de seguridad.

## Limitaciones y siguientes pasos
1. Validar etiquetas con especialistas y evaluar mensajes de plantillas nuevas o logs reales anonimizados.
2. Usar una prueba temporal o separada por plantillas para reducir el riesgo de fuga de información.
3. Añadir características estructuradas como duración, estado HTTP, reintentos y métricas del pool.
4. Medir el pipeline completo y el impacto en diagnóstico mediante un piloto en modo sombra.
5. Las estimaciones de ROI son un escenario de negocio basado en supuestos, no ahorros observados.
'''
(ROOT/'README.md').write_text(readme,encoding='utf-8')
(ROOT/'requirements.txt').write_text('pandas>=2.0\nnumpy>=1.24\nscikit-learn>=1.3\nmatplotlib>=3.7\nopenpyxl>=3.1\npython-docx>=1.1\npython-pptx>=0.6.23\njoblib>=1.3\n',encoding='utf-8')
print('DOCX, XLSX and PPTX created in',ROOT)
